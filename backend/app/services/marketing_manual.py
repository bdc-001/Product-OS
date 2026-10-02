"""Editable Google Docs source with semantic headings and preserved operator steps."""
from html import escape
from pathlib import Path
import base64
import logging
import math
import re

log = logging.getLogger(__name__)

GOOGLE_DOC = "application/vnd.google-apps.document"
FONT_DIR = Path(__file__).resolve().parents[1] / "fonts"
NUNITO_REGULAR = FONT_DIR / "Nunito-Regular.ttf"
NUNITO_BOLD = FONT_DIR / "Nunito-Bold.ttf"


def _logo_png() -> bytes:
    import fitz
    from app.services.brand import logo_data_uri
    uri = logo_data_uri(dark=False)
    kind = uri.split(';', 1)[0].rsplit('/', 1)[-1].replace('svg+xml', 'svg')
    with fitz.open(stream=base64.b64decode(uri.split(',', 1)[1]), filetype=kind) as logo:
        return logo[0].get_pixmap(matrix=fitz.Matrix(3, 3), alpha=True).tobytes('png')


def manual_header() -> str:
    image = base64.b64encode(_logo_png()).decode('ascii')
    return ('<div style="width:100%;text-align:right;padding:0 16mm;font-size:0">'
            f'<img src="data:image/png;base64,{image}" style="width:120px"></div>')


def _subheading(line: str, following: str = "") -> str:
    """Explicit Markdown subheadings, plus short labels from older saved drafts."""
    if line.startswith('### '):
        return line[4:].strip()
    if (following.strip() and 1 <= len(line.split()) <= 9
            and not re.search(r'[.!?:]$|^\d+[.)]\s|^[-•*]\s', line)
            and (len(following.split()) > 9 or re.match(r'^[-•*]\s', following))):
        return line
    return ""


def verify_google_doc(service, file_id: str, content: dict) -> None:
    """Read back native text after conversion; a successful upload alone is insufficient."""
    raw = service.files().export_media(fileId=file_id, mimeType="text/plain").execute(num_retries=2)
    actual = raw.decode("utf-8") if isinstance(raw, bytes) else str(raw)
    def normal(text):
        return " ".join(re.sub(r"(?m)^\s*(?:\d+[.)]|[-•*]|#{2,4})\s+", "", text).split())
    expected = [content["title"], content["dek"], content["cta"]]
    for section in content["sections"]:
        expected.append(section["body"])
        if section.get("role") != "overview":
            expected.append(section["heading"])
    expected.extend(figure['caption'] for figure in content.get('figures', []))
    if any(normal(part) not in normal(actual) for part in expected):
        raise RuntimeError("Google Docs conversion did not preserve the complete manual. Local copies are retained; retry delivery.")
    if content.get("figures"):
        import fitz
        pdf = service.files().export_media(fileId=file_id, mimeType="application/pdf").execute(num_retries=2)
        with fitz.open(stream=pdf, filetype="pdf") as document:
            images = {image[0] for page in document for image in page.get_images()}
            if len(images) < len(content["figures"]):
                raise RuntimeError("Google Docs conversion lost a feature image. Delivery needs repair.")
            fonts = [font[3].lower() for page in document for font in page.get_fonts()]
            if not any("nunito" in name for name in fonts):
                raise RuntimeError("Google Docs conversion lost Nunito typography. Delivery needs repair.")
            wrong_face = [span['text'] for page in document for block in page.get_text('dict')['blocks']
                          for line in block.get('lines', []) for span in line['spans']
                          if sum(char.isalpha() for char in span['text']) > 5 and 'nunito' not in span['font'].lower()]
            if wrong_face:
                raise RuntimeError("Google Docs body or headings use a fallback font instead of Nunito. Delivery needs repair.")


def manual_body(content: dict, folder: Path | None = None) -> str:
    parts = [f"<h1>{escape(content['title'])}</h1>"]
    if content.get('dek'):
        parts.append(f"<p>{escape(content['dek'])}</p>")
    placed = set()
    for section in content["sections"]:
        if section.get('role') != 'overview':
            parts.append(f"<h2>{escape(section['heading'])}</h2>")
        lines = [line.strip() for line in section["body"].splitlines() if line.strip()]
        list_type = None
        for line_index, line in enumerate(lines):
            line = line.strip()
            if not line:
                continue
            subhead = _subheading(line, lines[line_index + 1] if line_index + 1 < len(lines) else '')
            if subhead:
                if list_type:
                    parts.append(f"</{list_type}>")
                    list_type = None
                parts.append(f"<h3>{escape(subhead)}</h3>")
                continue
            step = re.match(r"^\d+[.)]\s+(.+)$", line)
            bullet = re.match(r"^[-•*]\s+(.+)$", line)
            new_type = 'ol' if step else 'ul' if bullet else None
            if new_type != list_type:
                if list_type:
                    parts.append(f"</{list_type}>")
                if new_type:
                    parts.append(f"<{new_type}>")
                list_type = new_type
            if step or bullet:
                parts.append(f"<li>{escape((step or bullet)[1])}</li>")
            else:
                parts.append(f"<p>{escape(line)}</p>")
        if list_type:
            parts.append(f"</{list_type}>")
        for index, figure in enumerate(content.get("figures", [])):
            if index in placed or figure["after_role"] != section.get("role") or folder is None:
                continue
            path = folder / f"release_notes-figure-{index + 1:02}.png"
            encoded = base64.b64encode(path.read_bytes()).decode()
            parts.append(f'<figure><img alt="{escape(figure["heading"])}" src="data:image/png;base64,{encoded}"><figcaption>{escape(figure["caption"])}</figcaption></figure>')
            placed.add(index)
    parts.append(f"<p class=\"manual-cta\">{escape(content['cta'])}</p>")
    return "\n".join(parts)


def _nunito_faces() -> str:
    faces = []
    for weight, path in ((400, NUNITO_REGULAR), (700, NUNITO_BOLD)):
        if not path.is_file():
            continue
        b64 = base64.b64encode(path.read_bytes()).decode("ascii")
        faces.append(
            f'@font-face{{font-family:"Nunito";src:url(data:font/ttf;base64,{b64}) format("truetype");'
            f"font-weight:{weight};font-display:block;}}"
        )
    return "\n".join(faces)


def render_manual_figures(content: dict, folder: Path) -> list[tuple[str, str]]:
    """Render source-grounded diagrams locally; no public image hosting is needed for Docs."""
    from playwright.sync_api import sync_playwright
    files = []
    if not content.get("figures"):
        return files
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        try:
            page = browser.new_page(viewport={"width": 1400, "height": 510}, device_scale_factor=1, java_script_enabled=False)
            page.route("**/*", lambda route: route.abort())
            for index, figure in enumerate(content["figures"]):
                items = []
                for i, item in enumerate(figure["items"]):
                    if i and figure["kind"] == "flow":
                        items.append('<span class="arrow">→</span>')
                    items.append(f'<div class="node"><b>{escape(item["lead"])}</b><p>{escape(item["text"])}</p></div>')
                source = ('<html><head><style>' + _nunito_faces() + '''
                *{box-sizing:border-box}body{margin:0;background:#eef4ff;color:#102342;font-family:Nunito}
                main{padding:40px 48px;width:1400px;min-height:510px;border:2px solid #d1e1fd}
                .label{font-size:20px;letter-spacing:1px;color:#1658e1}h1{font-size:40px;line-height:1.35;margin:12px 0 26px}
                .diagram{display:flex;align-items:stretch;gap:18px}.node{flex:1;min-width:0;min-height:228px;border:1px solid #afcafb;border-radius:18px;background:white;padding:28px}
                .node:last-child{background:#0a2e7a;color:white;border-color:#0a2e7a}b{display:block;font-size:34px;line-height:1.2}p{font-size:29px;line-height:1.3;margin:18px 0 0}.arrow{font-size:42px;color:#1a62f2;align-self:center}
                </style></head><body><main><div class="label">CONCEPTUAL WORKFLOW</div><h1>'''
                + escape(figure["heading"]) + '</h1><div class="diagram">' + ''.join(items) + '</div></main></body></html>')
                source = source.replace('CONCEPTUAL WORKFLOW', 'CONCEPTUAL ' + ('WORKFLOW' if figure['kind'] == 'flow' else 'COMPARISON'))
                page.set_content(source, wait_until="load")
                page.evaluate("() => document.fonts.ready")
                height = math.ceil(page.locator('main').bounding_box()['height'])
                if height > 720:
                    raise ValueError("Manual diagram needs shorter labels; its natural height exceeds 720px.")
                # Let an honest extra line increase the figure height. The manual
                # flows around images; squeezing or cropping text is unnecessary.
                page.set_viewport_size({"width": 1400, "height": height})
                overflow = page.evaluate("""limit => [...document.querySelectorAll('.node,h1')].map(el=>({text:el.innerText,scroll:el.scrollHeight,height:el.clientHeight,bottom:el.getBoundingClientRect().bottom})).filter(el=>el.scroll>el.height+1 || el.bottom>limit)""", height - 20)
                if overflow:
                    page.screenshot(path=str(folder / f"manual-diagram-{index + 1}-diagnostic.png"))
                    raise ValueError(f"Manual diagram exceeds its image: {overflow}")
                name = f"release_notes-figure-{index + 1:02}.png"
                page.screenshot(path=str(folder / name))
                files.append((name, "image/png"))
        finally:
            browser.close()
    return files


def manual_design(content: dict, folder: Path | None = None) -> str:
    """Nunito operator manual with labelled explanatory images in reading order."""
    return (
        '<!doctype html><html><head><meta charset="utf-8"><style>'
        + _nunito_faces()
        + "@page{size:A4;margin:22mm 16mm 16mm}"
        "*{box-sizing:border-box}"
        'body{margin:0;font-family:Nunito,Arial,sans-serif;font-size:11pt;line-height:1.45;color:#182b46}'
        "p.brand{margin:0 0 18px;font-size:9pt;letter-spacing:1.2px;color:#1a62f2;font-weight:700}"
        "h1{font-size:22pt;line-height:1.2;margin:0 0 12px;font-weight:700;letter-spacing:-.4px}"
        "h2{font-size:14pt;line-height:1.3;margin:22px 0 8px;color:#102342;font-weight:700;break-after:avoid}"
        "h3{font-size:11pt;line-height:1.4;margin:14px 0 5px;font-weight:700;break-after:avoid}"
        "h2+p,h3+p,p:has(+ul),p:has(+ol){break-after:avoid}"
        "p{margin:0 0 9px;orphans:3;widows:3}li{margin-bottom:8px;break-inside:avoid}"
        "ol,ul{padding-left:22px}h1+p{font-size:11pt;margin-bottom:18px}"
        "figure{margin:16px 0 20px;break-inside:avoid}figure img{width:100%;height:auto;display:block}"
        "figcaption{font-size:9pt;line-height:1.4;color:#435772;margin-top:8px}"
        ".manual-cta{border-top:1px solid #959595;padding-top:14px;margin-top:22px;font-style:italic;color:#6A6A6A}"
        "</style></head><body>"
        + manual_body(content, folder) + "</body></html>"
    )


def write_google_doc_source(content: dict, folder: Path) -> dict:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.71)
    section.left_margin = section.right_margin = Inches(.71)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Number", "List Bullet", "Caption"):
        style = document.styles[name]
        style.font.name = "Nunito"
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if 'theme' in attribute.lower():
                del fonts.attrib[attribute]
        for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + script), 'Nunito')
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string("182B46")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
    document.styles['Caption'].font.bold = False
    document.styles['Caption'].font.size = Pt(9)
    title_properties = document.styles['Title'].element.get_or_add_pPr()
    for border in list(title_properties.findall(qn('w:pBdr'))):
        title_properties.remove(border)
    for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 12)):
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.bold = True
        document.styles[name].paragraph_format.keep_with_next = True
        document.styles[name].paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
    # Rasterize the checked-in vector lockup for Word/Docs' native header image.
    import io
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run().add_picture(io.BytesIO(_logo_png()), width=Inches(1.25))
    section.header_distance = Inches(.25)
    document.add_paragraph(content["title"], "Title")
    if content.get('dek'):
        document.add_paragraph(content["dek"])
    placed = set()
    for section in content["sections"]:
        if section.get('role') != 'overview':
            document.add_heading(section["heading"], level=1)
        lines = [line.strip() for line in section["body"].splitlines() if line.strip()]
        for line_index, line in enumerate(lines):
            if not line.strip():
                continue
            subhead = _subheading(line, lines[line_index + 1] if line_index + 1 < len(lines) else '')
            if subhead:
                document.add_heading(subhead, level=2)
                continue
            step = re.match(r"^\d+[.)]\s+(.+)$", line.strip())
            bullet = re.match(r"^[-•*]\s+(.+)$", line.strip())
            paragraph = document.add_paragraph((step or bullet)[1] if step or bullet else line.strip(),
                                              "List Number" if step else "List Bullet" if bullet else "Normal")
            if line.endswith(':'):
                paragraph.paragraph_format.keep_with_next = True
        for index, figure in enumerate(content.get("figures", [])):
            if index in placed or figure["after_role"] != section.get("role"):
                continue
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            picture = paragraph.add_run().add_picture(str(folder / f"release_notes-figure-{index + 1:02}.png"), width=Inches(6.85))
            picture._inline.docPr.set("descr", figure["heading"])
            caption = document.add_paragraph(figure["caption"], "Caption")
            caption.paragraph_format.space_after = Pt(12)
            placed.add(index)
    cta = document.add_paragraph(content["cta"])
    cta.paragraph_format.space_before = Pt(16)
    for run in cta.runs:
        run.italic = True
        run.font.color.rgb = RGBColor.from_string('6A6A6A')
    name = "release_notes-google-doc.docx"
    document.save(folder / name)
    return {
        "filename": name,
        "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "channel": "release_notes",
        "drive_mime": GOOGLE_DOC,
        "drive_name": content["title"] + " - Release Note",
    }


def apply_nunito(doc_id: str) -> None:
    """Pin every named style and text run to Nunito after HTML-to-Docs conversion."""
    from app.services.gdocs import _docs, _docs_batch
    from app.services.gdocs_style import named_style_request, font_family, pt, rgb_color

    docs = _docs()
    docs.documents().batchUpdate(
        documentId=doc_id,
        body={"requests": [
            named_style_request("NORMAL_TEXT", size=11, bold=False, fg="050505", space_below=8, family="Nunito"),
            named_style_request("TITLE", size=22, bold=True, fg="102342", space_below=10, family="Nunito"),
            named_style_request("SUBTITLE", size=11, bold=False, fg="3F3F3F", space_below=8, family="Nunito"),
            named_style_request("HEADING_1", size=14, bold=True, fg="102342", space_below=10, family="Nunito"),
            named_style_request("HEADING_2", size=12, bold=True, fg="1A62F2", space_below=6, family="Nunito"),
            named_style_request("HEADING_3", size=12, bold=True, fg="102342", space_below=4, family="Nunito"),
        ]},
    ).execute()
    doc = docs.documents().get(documentId=doc_id).execute()
    body_content = (doc.get("body") or {}).get("content") or [{"endIndex": 2}]
    end = int(body_content[-1].get("endIndex") or 2)
    if end > 2:
        _docs_batch(doc_id, [{
            "updateTextStyle": {
                "range": {"startIndex": 1, "endIndex": end - 1},
                "textStyle": {"weightedFontFamily": font_family(False, "Nunito")},
                "fields": "weightedFontFamily",
            }
        }])
    requests = []
    heading_types = {"TITLE", "HEADING_1", "HEADING_2", "HEADING_3"}
    for element in (doc.get("body") or {}).get("content") or []:
        paragraph = element.get("paragraph")
        if not paragraph:
            continue
        named = (paragraph.get("paragraphStyle") or {}).get("namedStyleType")
        if named not in heading_types:
            continue
        start = int(element.get("startIndex") or 1)
        end_index = int(element.get("endIndex") or start)
        size = 22 if named == "TITLE" else 14 if named == "HEADING_1" else 12
        fg = "1A62F2" if named == "HEADING_2" else "102342"
        requests.append({
            "updateTextStyle": {
                "range": {"startIndex": start, "endIndex": max(start + 1, end_index - 1)},
                "textStyle": {
                    "bold": True,
                    "fontSize": pt(size),
                    "foregroundColor": rgb_color(fg),
                    "weightedFontFamily": font_family(True, "Nunito"),
                },
                "fields": "bold,fontSize,foregroundColor,weightedFontFamily",
            }
        })
    if requests:
        _docs_batch(doc_id, requests)
