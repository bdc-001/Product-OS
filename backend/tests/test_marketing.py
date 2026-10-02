from types import SimpleNamespace
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from app.database import Base, get_db
from app.models import CodebaseModule, CodebaseSnapshot, MarketingCampaign, MarketingFeature, MarketingScan
from app.services import marketing, marketing_manager as manager

@pytest.fixture(autouse=True)
def no_live_manager(monkeypatch, tmp_path):
    monkeypatch.setattr(marketing, "CAMPAIGN_DIR", tmp_path)
    monkeypatch.setattr(manager, "ROOT", tmp_path)
    monkeypatch.setattr(manager, "chat_json", lambda *a, **k: pytest.fail("Tests must mock PMM calls"))
    monkeypatch.setattr("app.config.settings.marketing_sheet_id", "")
from app.services.marketing_video import bind_scene, document_html

@pytest.fixture
def db(tmp_path):
    # Production uses file-backed SQLite; concurrent sessions must own connections.
    engine = create_engine(f"sqlite:///{tmp_path / 'marketing.sqlite'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    yield session
    session.close()
    engine.dispose()

@pytest.fixture
def client(db):
    from app.api.routers.marketing import router
    app = FastAPI(); app.include_router(router)
    app.dependency_overrides[get_db] = lambda: db
    return TestClient(app)

def seed(db, **kwargs):
    row = MarketingFeature(name="AI campaigns", identity="ai campaigns", description="Create campaigns", audience="Sales teams", **kwargs)
    db.add(row); db.commit()
    return row

def test_sheet_updates_are_validated_and_versioned(client, db):
    assert client.post("/marketing/features", json={"name": "  "}).status_code == 422
    created = client.post("/marketing/features", json={"name": "Campaign analytics", "description": "Insights"}).json()
    assert created["source"] == "manual"
    assert client.post("/marketing/features", json={"name": " Campaign ANALYTICS "}).status_code == 409
    saved = client.put(f"/marketing/features/{created['id']}", json={"name": "Campaign analytics", "description": "Edited", "revision": 1}).json()
    assert saved["revision"] == 2
    assert client.put(f"/marketing/features/{created['id']}", json={"name": "Campaign analytics", "revision": 1}).status_code == 409
    assert db.get(MarketingFeature, created["id"]).description == "Edited"

def test_generation_requires_useful_description(client, db):
    row = seed(db, status="draft")
    assert client.post("/marketing/generate", json={"feature_ids": [row.id]}).status_code == 400
    assert client.post("/marketing/generate", json={"feature_ids": []}).status_code == 422
    assert client.post("/marketing/generate", json={"feature_ids": [999]}).status_code == 404

def test_discovery_waits_for_first_manual_entry(db, monkeypatch):
    monkeypatch.setattr(marketing, "chat_json", lambda *a, **k: pytest.fail("No model call before initial sheet"))
    assert marketing.discover_features(db)["count"] == 0

def campaign_content():
    raw = {"strategy": {"positioning": "Manage campaigns in Voice"}, "video": {"title": "Voice campaigns", "scenes": []}}
    for channel in marketing.CHANNELS:
        raw[channel] = {"title": channel.title(), "body": "Complete feature content with a customer benefit. " * 8}
    for i in range(12):
        raw["video"]["scenes"].append({"kind": "hook" if i == 0 else "cta" if i == 11 else "benefit",
            "headline": "Understand your campaigns", "body": "See how your campaigns are performing.",
            "narration": "See how each campaign performs in one clear view.", "evidence": "User supplied description"})
    return raw

def test_content_validation_rejects_incomplete_outputs_and_overflow():
    raw = campaign_content()
    assert marketing.validate_content(raw) == raw
    raw["linkedin"]["body"] = "Outline"
    with pytest.raises(ValueError, match="incomplete"): marketing.validate_content(raw)
    raw = campaign_content(); raw["video"]["scenes"][0]["headline"] = "very " * 20
    with pytest.raises(ValueError, match="readable"): marketing.validate_content(raw)

def test_audio_timing_prevents_cutoffs_and_overlapping_captions():
    timed = bind_scene(campaign_content()["video"]["scenes"][0], 4.25, 150)
    assert timed["frames"] % 3 == 0 and timed["frames"] >= timed["audioFrames"] + 9
    assert timed["from"] == 150
    assert timed["captions"][0]["to"] <= timed["captions"][1]["from"]
    assert timed["captions"][-1]["to"] < timed["frames"]
    with pytest.raises(ValueError): bind_scene(campaign_content()["video"]["scenes"][0], float("nan"), 0)

def test_printable_documents_escape_model_html():
    rendered = document_html("<script>alert(1)</script>", "# Title\n\n<script>alert(2)</script>")
    assert "<script>" not in rendered and "&lt;script&gt;" in rendered and "<h2>Title</h2>" in rendered

def test_generation_failure_never_creates_placeholder_content(db, tmp_path, monkeypatch):
    row = seed(db, status="ready")
    campaign = MarketingCampaign(feature_id=row.id, feature_revision=1, feature_snapshot=marketing.feature_out(row))
    db.add(campaign); db.commit()
    monkeypatch.setattr(marketing, "CAMPAIGN_DIR", tmp_path)
    def fail(*a, **k): raise RuntimeError("Model unavailable")
    monkeypatch.setattr(manager, "chat_json", fail)
    assert marketing.produce_campaign(db, campaign)["ok"] is False
    assert campaign.status == "failed" and not campaign.content and not campaign.assets

def test_retry_preserves_content_and_skips_successful_video(db, tmp_path, monkeypatch):
    row = seed(db, status="ready")
    campaign = MarketingCampaign(feature_id=row.id, feature_revision=1, feature_snapshot=marketing.feature_out(row),
        content=campaign_content(), assets=[{"filename": "landscape.mp4"}, {"filename": "portrait.mp4"}], status="partial")
    db.add(campaign); db.commit()
    monkeypatch.setattr(marketing, "CAMPAIGN_DIR", tmp_path)
    monkeypatch.setattr(marketing, "chat_json", lambda *a, **k: pytest.fail("Must not regenerate saved content"))
    monkeypatch.setattr("app.services.marketing_video.render_campaign_video", lambda *a, **k: pytest.fail("Must not rerender completed videos"))
    uploaded = []
    monkeypatch.setattr(marketing, "upload_campaign", lambda db, row: uploaded.append(row.id))
    assert marketing.produce_campaign(db, campaign)["ok"] is True
    assert uploaded == [campaign.id] and campaign.status == "completed"

def test_upload_checkpoint_reuses_remote_files(db, tmp_path, monkeypatch):
    row = seed(db)
    campaign = MarketingCampaign(feature_id=row.id, feature_revision=1, feature_snapshot=marketing.feature_out(row), drive_folder_id="folder",
        assets=[{"filename": "one.pdf", "mime": "application/pdf", "file_id": "existing", "drive_root": "root"},
                {"filename": "two.pdf", "mime": "application/pdf"},
                {"filename": "notes.md", "mime": "text/markdown"}])
    db.add(campaign); db.commit()
    folder = tmp_path / str(campaign.id); folder.mkdir()
    (folder / "one.pdf").write_bytes(b"%PDF")
    (folder / "two.pdf").write_bytes(b"%PDF")
    updated = []
    class Files:
        def list(self, **kwargs):
            return SimpleNamespace(execute=lambda **k: {"files": [{"id": "recovered", "webViewLink": "https://drive.google.com/file/recovered"}]})
        def create(self, **kwargs): pytest.fail("Recovered uploads must not be duplicated")
        def update(self, **kwargs):
            updated.append(kwargs["fileId"])
            return SimpleNamespace(execute=lambda **k: {"id": kwargs["fileId"], "webViewLink": "https://drive.google.com/file/existing"})
    monkeypatch.setattr(marketing, "CAMPAIGN_DIR", tmp_path)
    monkeypatch.setattr(marketing, "drive_destination", lambda: {"configured": True, "folder_id": "root"})
    monkeypatch.setattr(marketing, "ensure_year_folder", lambda service, parent, name: f"{parent}-{name}")
    monkeypatch.setattr(marketing, "_service", lambda: SimpleNamespace(files=lambda: Files()))
    monkeypatch.setattr("app.services.marketing_manual.apply_nunito", lambda doc_id: None)
    marketing.upload_campaign(db, campaign)
    assert updated == []  # Delivered assets are immutable on a retry.
    assert [a.get("file_id") for a in campaign.assets] == ["existing", "recovered", None]

def test_download_only_allows_manifest_entries(client, db, tmp_path, monkeypatch):
    from app.api.routers import marketing as routes
    row = seed(db)
    campaign = MarketingCampaign(feature_id=row.id, feature_revision=1, feature_snapshot=marketing.feature_out(row), assets=[{"filename": "document.md", "mime": "text/markdown"}])
    db.add(campaign); db.commit()
    folder = tmp_path / str(campaign.id); folder.mkdir(); (folder / "document.md").write_text("safe")
    (folder / "private.txt").write_text("private")
    monkeypatch.setattr(routes, "CAMPAIGN_DIR", tmp_path)
    assert client.get(f"/marketing/campaigns/{campaign.id}/files/document.md").text == "safe"
    assert client.get(f"/marketing/campaigns/{campaign.id}/files/private.txt").status_code == 404

def test_csv_export_neutralizes_formulas(client, db):
    row = seed(db); row.script = '=HYPERLINK("x")'; db.commit()
    response = client.get("/marketing/features.csv")
    assert response.status_code == 200 and "'=HYPERLINK" in response.text
    assert "Module,Feature,Description,Use Case,Industry" in response.text


QUOTE = "View a breakdown of campaign results by channel."

def decision(verdict="approve", selected=("linkedin_post",)):
    selected = selected if verdict == "approve" else ()
    return {"verdict": verdict, "rationale": "A clear customer-facing capability with a specific buyer use case.",
        "audience": "Sales operations", "buyer_problem": "Teams need to understand their campaign results.",
        "positioning": "Understand campaign outcomes by channel.", "evidence_quotes": [QUOTE],
        "missing_evidence": ["Deployment availability is not documented."] if verdict == "defer" else [],
        "cta": "Explore how Voice supports campaign reporting.",
        "assignments": [{"format": key, "reason": "This format clearly explains the reporting workflow.", "angle": "Compare campaign outcomes by channel.",
                         "objective": "Help buyers understand the reporting capability.", "outline": ["Customer problem", "Supported workflow"]} for key in selected],
        "omitted_formats": {key: "Additional depth or visual treatment is not justified here." for key in manager.FORMATS if key not in selected}}


def mock_release(monkeypatch, raw):
    release = {"sha": "a" * 40, "base_sha": "b" * 40, "branch": "release/2026-09-10", "merged_at": "2026-09-10T10:00:00Z"}
    monkeypatch.setattr("app.services.release_detect.detect_release_merges", lambda **kw: [release])
    monkeypatch.setattr(manager, "release_batches", lambda release: [[{"path": "billing-service/internal/service/campaign.go", "text": QUOTE, "added_lines": [QUOTE]}]])
    monkeypatch.setattr(manager, "chat_json", lambda *a, **k: raw)
    return release


def discovered_item(verdict="approve"):
    return {"name": "Channel results", "description": QUOTE, "evidence_quote": QUOTE,
            "evidence_path": "billing-service/internal/service/campaign.go", "decision": decision(verdict)}


@pytest.mark.usefixtures("acme_workspace")
def test_release_discovery_adds_sheet_row_without_auto_queue(db, monkeypatch):
    original = seed(db, status="dismissed", notes="Preserve my notes")
    mock_release(monkeypatch, {"features": [discovered_item()]})
    assert marketing.discover_features(db)["count"] == 1
    assert marketing.discover_features(db)["count"] == 0
    assert db.query(MarketingCampaign).count() == 0
    feature = db.query(MarketingFeature).filter_by(name="Channel results").one()
    assert feature.status == "candidate" and feature.decision["verdict"] == "approve"
    assert feature.tag == "New" and feature.sheet_status == "Not started"
    assert feature.module == "Voice" and feature.summary
    assert feature.evidence[0]["sha"] == "a" * 40
    assert original.notes == "Preserve my notes" and original.status == "dismissed"

def test_release_failure_does_not_checkpoint_unverified_feature(db, monkeypatch):
    seed(db)
    item = discovered_item(); item["evidence_quote"] = "Invented evidence which is not in the diff."
    mock_release(monkeypatch, {"features": [item]})
    assert not marketing.discover_features(db)["ok"]
    assert db.query(MarketingScan).count() == 0
    assert db.query(MarketingCampaign).count() == 0
    monkeypatch.setattr(manager, "chat_json", lambda *a, **k: {"features": [discovered_item()]})
    assert marketing.discover_features(db)["count"] == 1


def test_discovery_rejects_unchanged_context_as_new_feature(db, monkeypatch):
    seed(db)
    mock_release(monkeypatch, {"features": [discovered_item()]})
    monkeypatch.setattr(manager, "release_batches", lambda r: [[{"path": "billing-service/internal/service/campaign.go", "text": QUOTE, "added_lines": ["internal refactor only"]}]])
    assert marketing.discover_features(db)["ok"] is False
    assert db.query(MarketingCampaign).count() == 0


def test_no_release_history_never_falls_back_to_unreleased_index(db, monkeypatch):
    seed(db)
    monkeypatch.setattr("app.services.release_detect.detect_release_merges", lambda **kw: None)
    assert not marketing.discover_features(db)["ok"]
    assert db.query(MarketingScan).count() == 0


def test_marketed_name_replaces_internal_names_and_evidence_still_matches():
    from app.services.marketing_content import apply_market_name, _source_span
    renamed = apply_market_name({
        "title": "Agent Analysis for operators",
        "body": "Open Agent Analysis/Monitor feature, then use Agent Monitor.",
        "customer_evidence": "Agent Analysis source quote stays",
    }, "Agent Monitoring")
    assert renamed["title"] == "Agent Monitoring for operators"
    assert renamed["body"] == "Open Agent Monitoring, then use Agent Monitoring."
    assert renamed["customer_evidence"] == "Agent Analysis source quote stays"
    source = "An operator opens Agent Analysis for one AI agent and reviews past analysis runs."
    assert _source_span([source], "An operator opens Agent Monitoring for one AI agent") == "An operator opens Agent Analysis for one AI agent"


def test_bind_keeps_source_quotes_and_drops_unsourced():
    source = (
        "A run evaluates recent interactions, groups root causes, and can suggest a prompt patch. "
        "The paid plan lets an operator start a run from the agent."
    )
    sourced = "A run evaluates recent interactions, groups root causes, and can suggest a prompt patch."
    benefit = "Operators can see why an agent is missing outcomes and update the prompt from the analysis instead of reading every call."
    paid = "The paid plan lets an operator start a run from the agent."
    assert manager.bind_evidence_quotes([sourced, benefit, paid], [source]) == [sourced, paid]
    with pytest.raises(ValueError):
        manager.bind_evidence_quotes([benefit], [source])


def test_pmm_decision_is_dynamic_and_evidence_bound():
    assert manager.validate_decision(decision(selected=("linkedin_carousel", "article")), QUOTE)["verdict"] == "approve"
    assert manager.validate_decision(decision("skip"), QUOTE)["assignments"] == []
    assert manager.validate_decision(decision("defer"), QUOTE)["missing_evidence"]
    bad = decision(); bad["assignments"].append(bad["assignments"][0])
    with pytest.raises(ValueError): manager.validate_decision(bad, QUOTE)
    bad = decision(); bad["evidence_quotes"] = ["Unsupported new capability in a made-up source"]
    with pytest.raises(ValueError): manager.validate_decision(bad, QUOTE)
    bad = decision(); bad["omitted_formats"] = {}
    with pytest.raises(ValueError): manager.validate_decision(bad, QUOTE)


def queued_campaign(db, verdict="approve", selected=("linkedin_post",)):
    feature = seed(db, status="candidate")
    feature.description = QUOTE
    db.commit()
    campaign = manager.queue_feature(db, feature, manager.validate_decision(decision(verdict, selected), QUOTE))
    db.commit()
    return campaign


@pytest.mark.parametrize("verdict,status", [("skip", "skipped"), ("defer", "deferred")])
def test_pmm_can_choose_no_production(db, tmp_path, monkeypatch, verdict, status):
    campaign = queued_campaign(db, verdict)
    monkeypatch.setattr(marketing, "CAMPAIGN_DIR", tmp_path)
    monkeypatch.setattr("app.services.marketing_content.write_format", lambda *a: pytest.fail("Unselected formats must never run"))
    monkeypatch.setattr(marketing, "upload_campaign", lambda *a: pytest.fail("No production to upload"))
    assert marketing.produce_campaign(db, campaign)["ok"]
    assert campaign.status == status
    assert (tmp_path / str(campaign.id) / "pmm-decision.json").is_file()


def test_failed_format_does_not_block_others_and_retry_is_resumable(db, tmp_path, monkeypatch):
    campaign = queued_campaign(db, selected=("linkedin_post", "article"))
    monkeypatch.setattr(marketing, "CAMPAIGN_DIR", tmp_path)
    writes, renders, uploads = [], [], []
    def write(brief, plan, assignment):
        writes.append(assignment["format"])
        return {"title": "Channel outcomes", "body": QUOTE}
    def render(key, content, folder, step):
        renders.append(key)
        if key == "linkedin_post" and renders.count(key) == 1:
            raise RuntimeError("Renderer disconnected")
        return [{"filename": f"{key}.md", "mime": "text/markdown", "channel": key}]
    monkeypatch.setattr("app.services.marketing_content.write_format", write)
    monkeypatch.setattr("app.services.marketing_content.render_format", render)
    monkeypatch.setattr(marketing, "upload_campaign", lambda db, c: uploads.append([a["filename"] for a in c.assets]))
    assert not marketing.produce_campaign(db, campaign)["ok"]
    assert "article.md" in uploads[0] and campaign.status == "partial"
    assert marketing.produce_campaign(db, campaign)["ok"]
    assert sorted(writes) == ["article", "linkedin_post"]
    assert sorted(renders) == ["article", "linkedin_post", "linkedin_post"]
    assert campaign.status == "completed"


def test_campaign_key_deduplicates_repeated_refresh_and_allows_edited_version(db):
    campaign = queued_campaign(db)
    feature = db.get(MarketingFeature, campaign.feature_id)
    assert manager.queue_feature(db, feature).id == campaign.id
    feature.revision += 1; db.commit()
    newer = manager.queue_feature(db, feature); db.commit()
    assert newer.id != campaign.id


def test_formats_overlap_and_checkpoint_on_coordinator_thread(db, monkeypatch):
    import threading
    barrier = threading.Barrier(2, timeout=3)
    coordinator = threading.get_ident()
    campaign = queued_campaign(db, selected=("linkedin_post", "article"))
    uploads = []
    def write(brief, plan, assignment):
        assert threading.get_ident() != coordinator
        barrier.wait()  # Serial execution fails this test.
        return {"title": "Specific channel outcomes", "body": QUOTE}
    def upload(session, row):
        assert threading.get_ident() == coordinator
        uploads.append({a.get("channel") for a in row.assets})
    monkeypatch.setattr("app.services.marketing_content.write_format", write)
    monkeypatch.setattr("app.services.marketing_content.render_format", lambda key, *args: [
        {"filename": key + ".pdf", "mime": "application/pdf", "channel": key}])
    monkeypatch.setattr(marketing, "upload_campaign", upload)
    assert manager.produce(db, campaign)["ok"]
    assert all(s["status"] == "completed" for s in campaign.content["formats"].values())
    assert len(uploads[0] - {"strategy"}) == 1
    assert {"linkedin_post", "article"} <= uploads[-1]


def test_multiple_selected_formats_share_one_pmm_campaign(client, db, monkeypatch):
    feature = seed(db); feature.description = QUOTE; db.commit()
    monkeypatch.setattr("app.api.routers.marketing.enqueue", lambda *args: None)
    response = client.post("/marketing/generate", json={"feature_ids": [feature.id], "formats": ["article", "feature_brief", "feature_image"]})
    assert response.status_code == 202
    rows = db.query(MarketingCampaign).all()
    assert len(rows) == 1
    assert set(rows[0].content["requested_formats"]) == {"article", "feature_brief", "feature_image"}


def test_queue_does_not_duplicate_formats_owned_by_an_active_campaign(db):
    feature = seed(db)
    first = manager.queue_feature(db, feature, requested_formats=["article", "feature_brief"])
    db.commit()
    assert manager.queue_feature(db, feature, requested_format="article").id == first.id
    second = manager.queue_feature(db, feature, requested_formats=["feature_brief", "linkedin_post"])
    db.commit()
    assert second.content["requested_formats"] == ["linkedin_post"]


def test_completed_selection_does_not_restart_failed_siblings(client, db):
    campaign = queued_campaign(db, selected=("linkedin_post", "article"))
    campaign.status = "partial"
    campaign.content = {**campaign.content, "formats": {
        "linkedin_post": {"status": "completed"}, "article": {"status": "failed"}}}
    db.commit()
    response = client.post("/marketing/generate", json={
        "feature_ids": [campaign.feature_id], "formats": ["linkedin_post"]})
    assert response.status_code == 409
    db.refresh(campaign)
    assert campaign.status == "partial"


def test_completed_selection_across_campaigns_does_not_restart_other_formats(client, db):
    campaign = queued_campaign(db, selected=("linkedin_post", "article"))
    feature = db.get(MarketingFeature, campaign.feature_id)
    campaign.status = "partial"
    campaign.content = {**campaign.content, "formats": {
        "linkedin_post": {"status": "completed"}, "article": {"status": "failed"}}}
    db.commit()
    other = manager.queue_feature(db, feature, requested_format="feature_brief")
    other.status = "completed"
    other.content = {**other.content, "formats": {"feature_brief": {"status":"completed"}}}
    db.commit()
    response = client.post("/marketing/generate", json={"feature_ids":[feature.id], "formats":["linkedin_post","feature_brief"]})
    assert response.status_code == 409
    db.refresh(campaign)
    assert campaign.status == "partial" and db.query(MarketingCampaign).count() == 2


def test_feature_stays_in_progress_until_sibling_campaign_finishes(db, monkeypatch):
    campaign = queued_campaign(db)
    feature = db.get(MarketingFeature, campaign.feature_id)
    sibling = manager.queue_feature(db, feature, requested_format="article")
    db.commit()
    monkeypatch.setattr("app.services.marketing_content.write_format", lambda *args: {"title": "Approved", "body": QUOTE})
    monkeypatch.setattr("app.services.marketing_content.render_format", lambda *args: [])
    monkeypatch.setattr(marketing, "upload_campaign", lambda *args: None)
    assert manager.produce(db, campaign)["ok"]
    assert sibling.status == "queued"
    assert feature.sheet_status == "In Progress"


def test_two_campaigns_run_together_with_separate_sessions(db, monkeypatch):
    import threading
    barrier = threading.Barrier(2, timeout=3)
    campaign = queued_campaign(db)
    feature = db.get(MarketingFeature, campaign.feature_id)
    feature.revision += 1; db.commit()
    second = manager.queue_feature(db, feature); db.commit()
    def finish(session, row, progress):
        assert session is not db
        barrier.wait()
        row.status = "completed"; session.commit()
        progress("done", {"status": "done"})
        return {"ok": True}
    monkeypatch.setattr(marketing, "produce_campaign", finish)
    result = manager.drain_queue(db)
    assert result["ok"] and len(result["campaigns"]) == 2
    assert campaign.status == second.status == "completed"


def test_static_design_failure_keeps_approved_copy(db, monkeypatch):
    from app.services.marketing_design import DesignQualityError
    campaign = queued_campaign(db)
    calls = []
    monkeypatch.setattr("app.services.marketing_content.write_format", lambda *args: calls.append(1) or {"title": "Approved title", "body": QUOTE})
    def fail(*args):
        raise DesignQualityError("Unreadable labels")
    monkeypatch.setattr("app.services.marketing_content.render_format", fail)
    monkeypatch.setattr(marketing, "upload_campaign", lambda *args: None)
    assert not manager.produce(db, campaign)["ok"]
    assert not manager.produce(db, campaign)["ok"]
    assert calls == [1]
    assert campaign.content["formats"]["linkedin_post"]["content"]["title"] == "Approved title"


def test_native_manual_upload_and_post_are_deliverables(db, tmp_path, monkeypatch):
    from app.services.marketing_manual import write_google_doc_source, GOOGLE_DOC
    monkeypatch.setattr("app.services.marketing_manual.apply_nunito", lambda doc_id: None)
    campaign = queued_campaign(db)
    folder = tmp_path / str(campaign.id); folder.mkdir()
    manual = {"title": "Agent Monitoring", "dek": "Review recent runs.",
        "sections": [{"heading": "Start a review", "body": "1. Open the agent.\n2. Review the results."}],
        "cta": "Review a run."}
    asset = write_google_doc_source(manual, folder)
    campaign.content = {**campaign.content, "formats": {"release_notes": {"content": manual}}}
    (folder / "linkedin_post.txt").write_text("Approved post")
    campaign.assets = [asset, {"filename": "linkedin_post.txt", "mime": "text/plain", "channel": "linkedin_post"}]
    db.commit()
    bodies = []
    class Files:
        def list(self, **kwargs): return SimpleNamespace(execute=lambda **k: {"files": []})
        def export_media(self, **kwargs):
            return SimpleNamespace(execute=lambda **k: b"Agent Monitoring\nReview recent runs.\nStart a review\n1. Open the agent.\n2. Review the results.\nReview a run.")
        def create(self, **kwargs):
            bodies.append(kwargs["body"])
            return SimpleNamespace(execute=lambda **k: {"id": str(len(bodies)), "webViewLink": "https://docs.google.com/document/d/1/edit"})
    monkeypatch.setattr(marketing, "drive_destination", lambda: {"configured": True, "folder_id": "root"})
    monkeypatch.setattr(marketing, "ensure_year_folder", lambda *args: "folder")
    monkeypatch.setattr(marketing, "_service", lambda: SimpleNamespace(files=lambda: Files()))
    marketing.upload_campaign(db, campaign)
    assert bodies[0]["mimeType"] == GOOGLE_DOC
    assert bodies[1]["name"] == "linkedin_post.txt"
    from docx import Document
    document = Document(folder / asset["filename"])
    assert [p.text for p in document.paragraphs if p.style.name == "List Number"] == ["Open the agent.", "Review the results."]
    assert document.styles["Normal"].font.name == "Nunito"
    marketing.upload_campaign(db, campaign)
    assert len(bodies) == 2


def test_dismissed_feature_cancels_pending_production(db, tmp_path, monkeypatch):
    campaign = queued_campaign(db)
    db.get(MarketingFeature, campaign.feature_id).status = "dismissed"; db.commit()
    monkeypatch.setattr(marketing, "CAMPAIGN_DIR", tmp_path)
    assert marketing.produce_campaign(db, campaign)["ok"]
    assert campaign.status == "cancelled"


def test_unattended_refresh_does_not_auto_start_production(db, monkeypatch):
    monkeypatch.setattr(manager, "discover_released_features", lambda *a: {"ok": True, "count": 1})
    called = []
    monkeypatch.setattr(manager, "drain_queue", lambda *a: called.append("drain") or {"ok": True})
    monkeypatch.setattr(manager, "schedule_pending", lambda *a: called.append("dispatch") or SimpleNamespace(id=9))
    result = manager.refresh_marketing(db)
    assert result["production"]["ok"] and result["production_job_id"] is None
    result = manager.refresh_marketing(db, lambda *a: None)
    assert result["production_job_id"] is None
    assert called == []


def test_refresh_marketing_syncs_sheet_then_discovers_then_writes_back(db, monkeypatch):
    seed(db)
    order = []
    monkeypatch.setattr(marketing, "seed_mastersheet", lambda db: order.append("seed") or {"ok": True, "added": 0})
    monkeypatch.setattr("app.services.gsheets.pull_into", lambda db: order.append("pull") or {"ok": True, "added": 1, "updated": 0})
    monkeypatch.setattr(manager, "discover_released_features", lambda *a: order.append("discover") or {"ok": True, "count": 1, "error": ""})
    monkeypatch.setattr("app.services.gsheets.push_from", lambda db: order.append("push") or {"ok": True, "count": 2})
    result = manager.refresh_marketing(db)
    assert order == ["seed", "pull", "discover", "push"]
    assert result["ok"] and result["count"] == 1
    assert result["filtered"]["hidden"] == 0
    assert result["sheet_pull"]["added"] == 1 and result["sheet_push"]["count"] == 2
    assert result["production_job_id"] is None


def test_queue_resumes_interrupted_stages_only(db, tmp_path, monkeypatch):
    campaign = queued_campaign(db)
    campaign.status = "rendering"; db.commit()
    monkeypatch.setattr(marketing, "CAMPAIGN_DIR", tmp_path)
    def finish(db, row, step):
        row.status = "completed"; db.commit()
        return {"ok": True}
    monkeypatch.setattr(marketing, "produce_campaign", finish)
    assert manager.drain_queue(db)["ok"]
    assert campaign.status == "completed"
    assert manager.drain_queue(db)["campaigns"] == []


def test_manager_and_specialists_explicitly_use_opus():
    from app.clients.llm import model_for
    assert model_for("marketing_manager") == model_for("marketing_writer") == model_for("marketing_review") == "claude-opus-5-5"


def test_specialist_repairs_editorial_failures(monkeypatch):
    from app.services import marketing_content as content
    calls = []
    def model(surface, prompt, payload, **kw):
        calls.append(surface)
        if surface == "marketing_writer":
            return {"title": "See campaign outcomes by channel", "body": QUALITY_POST}
        if calls.count("marketing_review") == 1:
            return {"passed": False, "issues": ["Clarify the buyer problem"], "scores": {}}
        return quality_review(payload["content"])
    monkeypatch.setattr(content, "chat_json", model)
    result = content.write_format({"description": QUOTE}, decision(), decision()["assignments"][0])
    assert result["review"]["passed"]
    assert calls == ["marketing_writer", "marketing_review", "marketing_writer", "marketing_review"]


def test_video_validation_never_invents_or_truncates_copy():
    from app.services.marketing_content import normalize_video
    raw = {"scenes": [{"headline": "This deliberately long headline needs a rewrite from the writer",
                       "narration": "First, check the approved source before deciding what to say in this scene.",
                       "visual": "steps", "labels": ["Review the source", "Choose a response"],
                       "label_cues": ["wrong cue", "another wrong cue"]}]}
    assert normalize_video(raw) == raw


def test_manual_requires_operator_workflow_and_boundaries():
    from app.services.marketing_content import ReleaseNotes
    with pytest.raises(ValueError, match="manual needs"):
        ReleaseNotes.model_validate({"title": "Agent Monitoring", "dek": "A detailed guide to reviewing an agent.",
            "sections": [{"role": "overview", "heading": "Overview", "body": QUOTE * 2}] * 6,
            "cta": "Review a recent run."})


def test_manual_sheet_rows_are_not_auto_queued_on_refresh(db, monkeypatch):
    feature = seed(db); feature.description = QUOTE; db.commit()
    monkeypatch.setattr(manager, "discover_released_features", lambda *a: {"ok": True, "count": 0})
    monkeypatch.setattr(manager, "schedule_pending", lambda db: SimpleNamespace(id=7))
    manager.refresh_marketing(db, lambda *a: None)
    manager.refresh_marketing(db, lambda *a: None)
    assert db.query(MarketingCampaign).count() == 0


def test_buyer_sellable_keeps_operator_tools_and_drops_engineer_work():
    assert marketing.buyer_sellable(
        "Phone Number rotation",
        summary="Rotate outbound DIDs so campaigns do not burn a single number.",
        description="Campaign operators spread outbound traffic across a pool of numbers.",
        audience="Collections and sales campaign operators",
    )[0]
    assert marketing.buyer_sellable(
        "Agent Testing", summary="Ship with confidence.",
        description="Try your agent on a real call before customers hear it.",
        audience="All 4 verticals",
    )[0]
    assert marketing.buyer_sellable(
        "AI Agent Builder", summary="Idea to live agent, no engineering.",
        description="Build a launch-ready voice agent without writing code.",
        audience="All 4 verticals", benefit="Idea to live agent, no engineering",
    )[0]
    assert not marketing.buyer_sellable(
        "Developer API", description="Engineering teams add leads programmatically.",
        audience="E-commerce, BFSI",
    )[0]
    assert not marketing.buyer_sellable(
        "Gemini Explicit Caching", description="Cache Gemini prompts for cheaper completions.",
        audience="Internal engineering",
    )[0]
    assert not marketing.buyer_sellable(
        "SSO & Tenancy Management", description="Employees log in with company SSO.",
        audience="All 4 verticals",
    )[0]
    assert not marketing.buyer_sellable(
        "Roles & Permissions", description="Control who can launch campaigns.",
        audience="All 4 verticals",
    )[0]
    assert marketing.buyer_sellable(
        "Warm Transfer Agent Whisper",
        summary="Coach the human before they join the live call.",
        description="Operators whisper context to the receiving agent during warm transfer.",
        audience="Contact-center and CX operations leaders",
        module="Released",
    )[0]
    assert marketing.buyer_sellable(
        "Campaign Scheduling",
        summary="Hold a campaign until Monday 10:00.",
        description="Ops schedules a collections blast for the next working window.",
        audience="BFSI, Collections",
    )[0]
    assert marketing.buyer_sellable(
        "Phone Number Round-Robin",
        summary="Rotate outbound DIDs so campaigns do not burn a single number.",
        description="Campaign operators spread outbound traffic across a pool of numbers.",
        audience="Collections and sales campaign operators",
    )[0]


def test_hide_unsellable_dismisses_without_deleting(db):
    keep = seed(db)
    drop = MarketingFeature(
        name="Developer API", identity="developer api",
        description="Engineering teams add, update, archive, and fetch leads programmatically.",
        audience="E-commerce, BFSI", status="draft", source="mastersheet",
    )
    db.add(drop)
    db.commit()
    result = marketing.hide_unsellable_features(db)
    assert result["hidden"] == 1 and result["names"] == ["Developer API"]
    assert db.get(MarketingFeature, drop.id).status == "dismissed"
    assert db.get(MarketingFeature, drop.id).tag == "Not selected"
    assert db.get(MarketingFeature, keep.id).status == "draft"
    assert db.query(MarketingFeature).count() == 2


def test_discovery_skips_engineer_facing_even_if_pmm_approves(db, monkeypatch):
    seed(db)
    item = discovered_item()
    item["name"] = "Gemini Explicit Caching"
    item["audience"] = "Internal engineering"
    item["decision"] = {**item["decision"], "audience": "Internal engineering"}
    mock_release(monkeypatch, {"features": [item]})
    result = marketing.discover_features(db)
    assert result["ok"] and result["count"] == 0 and result["rejected"] == 1
    assert db.query(MarketingFeature).filter_by(name="Gemini Explicit Caching").first() is None


def test_discovery_keeps_phone_number_rotation(db, monkeypatch):
    seed(db)
    item = discovered_item()
    item["name"] = "Phone Number rotation"
    item["summary"] = "Rotate outbound DIDs across campaigns."
    item["audience"] = "Campaign operators"
    item["decision"] = {**item["decision"], "audience": "Campaign operators"}
    mock_release(monkeypatch, {"features": [item]})
    assert marketing.discover_features(db)["count"] == 1
    row = db.query(MarketingFeature).filter_by(name="Phone Number rotation").one()
    assert row.status == "candidate" and row.decision["verdict"] == "approve"


def test_mastersheet_seed_imports_features(db, tmp_path):
    csv_path = tmp_path / "mastersheet.csv"
    csv_path.write_text(
        "Module,Feature,Priority,Hook,Use Case,Who To Market To,Start Date,End Date,Status,Script,Video\n"
        "AI Agents,Agent Testing,P1,Ship with confidence,Try your agent on a real call before customers.,All verticals,10-09-26,17-09-26,In Progress,,\n"
        "Platform,PII Masking,P1,Privacy built in,Sensitive data is masked across views.,BFSI,,,,,\n"
        "Integrations,Developer API,P3,Build exactly what your ops need,Engineering teams fetch leads programmatically.,E-commerce,,,,,\n",
        encoding="utf-8",
    )
    first = marketing.seed_mastersheet(db, csv_path)
    second = marketing.seed_mastersheet(db, csv_path)
    assert first["added"] == 2 and second["added"] == 0
    assert db.query(MarketingFeature).filter_by(name="Developer API").first() is None
    rows = db.query(MarketingFeature).order_by(MarketingFeature.name).all()
    assert rows[0].tag == "New" and rows[0].module == "AI Agents"
    assert rows[0].sheet_status == "In Progress"
    assert "real call" in rows[0].description
    assert rows[0].summary


def test_api_accepts_draft_for_pmm_judgment_without_human_ready(client, db, monkeypatch):
    from app.api.routers import marketing as routes
    feature = seed(db, status="draft"); feature.description = QUOTE; db.commit()
    monkeypatch.setattr(routes, "enqueue", lambda *a: {"id": 9, "kind": "marketing", "status": "queued"})
    monkeypatch.setattr(routes, "job_out", lambda row: row)
    assert client.post("/marketing/generate", json={"feature_ids": [feature.id]}).status_code == 202
    assert client.post("/marketing/generate", json={"feature_ids": [feature.id]}).status_code == 202
    assert db.query(MarketingCampaign).count() == 1
    assert db.get(MarketingFeature, feature.id).tag == "In pipeline"


def test_generate_queues_a_single_requested_content_type(client, db, monkeypatch):
    from app.api.routers import marketing as routes
    from app.services.marketing_status import content_inventory
    feature = seed(db, status="draft")
    feature.description = QUOTE
    db.commit()
    monkeypatch.setattr(routes, "enqueue", lambda *a: {"id": 11, "kind": "marketing", "status": "queued"})
    monkeypatch.setattr(routes, "job_out", lambda row: row)
    assert client.post("/marketing/generate", json={"feature_ids": [feature.id], "formats": ["linkedin_carousel"]}).status_code == 202
    campaign = db.query(MarketingCampaign).one()
    assert campaign.run_key == f"pmm-format-v1:{feature.id}:{feature.revision}:linkedin_carousel"
    assert campaign.content["requested_formats"] == ["linkedin_carousel"]
    assert client.post("/marketing/generate", json={"feature_ids": [feature.id], "formats": ["release_notes"]}).status_code == 202
    assert db.query(MarketingCampaign).count() == 2
    assert client.post("/marketing/generate", json={"feature_ids": [feature.id], "formats": ["not_a_format"]}).status_code == 422
    inventory = content_inventory(feature, db.query(MarketingCampaign).all())
    assert inventory["linkedin_carousel"]["status"] in {"queued", "planning", "generating", "rendering", "uploading"}
    assert inventory["linkedin_carousel"]["campaign_id"] == campaign.id
    assert inventory["article"]["status"] == "not_started"


def test_release_batches_pin_source_and_exclude_secrets(monkeypatch):
    calls = []
    def git(root, *args, **kwargs):
        calls.append(args)
        if "--name-only" in args:
            return SimpleNamespace(returncode=0, stdout="billing-service/internal/service/campaign.go\nbilling-service/.env\nother-product/README.md\n")
        return SimpleNamespace(returncode=0, stdout="+++ b/billing-service/internal/service/campaign.go\n unchanged context\n+" + QUOTE + "\n-deleted capability\n")
    monkeypatch.setattr("app.services.codebase._git", git)
    batches = list(manager.release_batches({"sha": "a" * 40, "base_sha": "b" * 40}))
    assert len(calls) == 2
    assert "a" * 40 in calls[1] and "b" * 40 in calls[1]
    assert batches[0][0]["added_lines"] == [QUOTE]
    assert "deleted capability" not in batches[0][0]["text"]


def test_refresh_fetches_branches_before_pmm_and_preserves_other_steps(db, monkeypatch, tmp_path):
    from app.services import refresh
    calls = []
    monkeypatch.setattr(refresh, "run_pipeline", lambda *a, **k: SimpleNamespace(status="completed", error="", jira_count=1, cliq_count=1))
    monkeypatch.setattr(refresh, "sync_roadmap_statuses", lambda *a, **k: None)
    monkeypatch.setattr(refresh, "codebase_root", lambda: tmp_path)
    monkeypatch.setattr(refresh, "latest_snapshot", lambda db: None)
    monkeypatch.setattr(refresh, "update_index", lambda *a, **k: calls.append("index") or {})
    monkeypatch.setattr(refresh, "list_recent_branches", lambda **k: calls.append("fetch") or {"branches": []})
    monkeypatch.setattr("app.services.release_worker.detect_and_enqueue", lambda *a, **k: calls.append("releases") or [])
    monkeypatch.setattr("app.services.competitive.refresh_all", lambda *a, **k: calls.append("competitors") or {"ok": True, "created": 2})
    monkeypatch.setattr(manager, "refresh_marketing", lambda *a, **k: calls.append("marketing") or {"ok": True})
    result = refresh.refresh_platform(db, lambda *a: None)
    assert result["ok"] and result["steps"]["marketing"]["ok"]
    assert result["steps"]["competitors"]["count"] == 2
    assert calls == ["index", "fetch", "releases", "competitors", "marketing"]


def test_layout_failure_requests_revised_copy_on_retry(db, tmp_path, monkeypatch):
    campaign = queued_campaign(db)
    monkeypatch.setattr(marketing, "CAMPAIGN_DIR", tmp_path)
    calls = []
    def write(*args, **kwargs):
        calls.append(kwargs)
        return {"title": "Feature report", "body": QUOTE}
    attempts = []
    def render(*args):
        attempts.append(1)
        if len(attempts) == 1: raise ValueError("Headline exceeds its safe area")
        return []
    monkeypatch.setattr("app.services.marketing_content.write_format", write)
    monkeypatch.setattr("app.services.marketing_content.render_format", render)
    monkeypatch.setattr(marketing, "upload_campaign", lambda *a: None)
    assert marketing.produce_campaign(db, campaign)["ok"]
    assert len(attempts) == 2
    assert calls[1]["repair"] == "Headline exceeds its safe area"


def visual_brief(**overrides):
    return {
        "product": "Analytics",
        "feature": "Channel Breakdown",
        "title": "See which channel carries each campaign",
        "dek": "Channel Breakdown splits every campaign result by calling, WhatsApp, SMS and email, so reviews start from the channel that moved instead of one blended total.",
        "audience": "Campaign managers and operations leads who review outbound campaigns",
        "positioning": {"problem": "Campaign reviews start from one blended total, so teams cannot tell which channel moved results.",
                        "buyer": "Heads of campaign operations; used weekly by campaign managers.",
                        "why_pay": "Budget and effort go to the channel that actually works instead of being spread on guesswork."},
        "problem": {"heading": "One total hides the channel that moved",
                    "text": "A campaign total shows how much happened, not which channel did the work or which one needs attention."},
        "solution": {"heading": "Every result, split by channel",
                     "text": "Open a campaign and read calling, WhatsApp, SMS and email side by side for the same period and filters."},
        "visuals": [{"kind": "diagram", "heading": "How a channel review runs", "diagram": {"kind": "flow", "nodes": [
            {"label": "Open the campaign", "text": "Pick the campaign under review."},
            {"label": "Read by channel", "text": "Compare channels in the same window."},
            {"label": "Carry one question", "text": "Bring that difference into the review."}]}}],
        "value": [
            {"icon": "layers", "title": "Spot the channel that moved", "text": "Each channel keeps its own row, so a drop is easy to place."},
            {"icon": "filter", "title": "Filters that carry over", "text": "The breakdown uses the period and conditions you already applied."},
            {"icon": "chart", "title": "Totals that reconcile", "text": "Channel rows add up to the campaign total on the same card."}],
        "cta": "Available in Acme Voice Analytics. Ask your Acme team for a walkthrough.",
        **overrides}


def visual_carousel(**slide_overrides):
    slides = [
        {"layout": "hook", "visual": "statement", "headline": "One campaign total hides the answer",
         "body": "The number tells you how much happened.", "points": [], "evidence": QUOTE},
        {"layout": "problem", "visual": "contrast", "headline": "Totals against channels",
         "body": "Same campaign, two very different readings.", "points": ["One blended total", "Four channel results"], "evidence": QUOTE},
        {"layout": "steps", "visual": "steps", "headline": "How the review runs",
         "body": "Three moves before the meeting.", "points": ["Open the campaign", "Read by channel", "Carry one question"], "evidence": QUOTE},
        {"layout": "proof", "visual": "spotlight", "headline": "Voice breaks results out by channel",
         "body": "The breakdown sits with the campaign.", "points": ["Per channel results"], "evidence": QUOTE},
        {"layout": "benefit", "visual": "spotlight", "headline": "Reviews start on the difference",
         "body": "Discussion opens where results diverge.", "points": ["A focused question"], "evidence": QUOTE},
        {"layout": "cta", "visual": "statement", "headline": "Read your last campaign again",
         "body": "Start with the channel that moved.", "points": [], "evidence": QUOTE},
    ]
    if slide_overrides:
        slides[1] = {**slides[1], **slide_overrides}
    return {"title": "Read a campaign channel by channel", "body": QUALITY_POST, "slides": slides}


def test_brief_is_a_sales_one_pager_not_a_prose_document():
    from app.services.marketing_content import validate_format
    value = validate_format("feature_brief", visual_brief())
    assert "## One total hides the channel that moved" in value["body"] and "## Why teams pay for it" in value["body"]
    assert "guesswork" not in value["body"]
    assert 120 <= len(value["body"].split()) <= 260
    long_card = {"icon": "chart", "title": "Totals that reconcile", "text": "Channel rows add up to the campaign total on the same card. " * 2}
    with pytest.raises(ValueError, match="at most|within 16 words"):
        validate_format("feature_brief", visual_brief(value=[*visual_brief()["value"][:2], long_card]))
    with pytest.raises(ValueError, match="sells the outcome"):
        validate_format("feature_brief", visual_brief(title="Channel Breakdown"))
    with pytest.raises(ValueError, match="positioning"):
        validate_format("feature_brief", visual_brief(positioning=None))


def test_display_copy_rejects_internal_component_names():
    from app.services.marketing_content import validate_format
    raw = visual_brief()
    raw['value'][0]['text'] = 'Open PromptDiffModal to review the proposed change.'
    with pytest.raises(ValueError, match='internal code identifiers'):
        validate_format('feature_brief', raw)


def test_manual_docx_preserves_images_nunito_and_semantic_structure(tmp_path):
    from app.services.marketing_manual import write_google_doc_source
    from docx import Document
    from docx.oxml.ns import qn
    from PIL import Image
    for index, color in enumerate(('blue', 'navy'), 1):
        Image.new('RGB', (140, 51), color).save(tmp_path / f'release_notes-figure-{index:02}.png')
    content = {'title': 'Feature operator manual', 'dek': 'Review the results.', 'cta': 'Review a run.',
               'sections': [{'role': 'overview', 'heading': 'Overview', 'body': 'Review the findings.'},
                            {'role': 'workflow', 'heading': 'Workflow', 'body': '1. Open a run.\n2. Review findings.'},
                            {'role': 'action', 'heading': 'Action', 'body': '- Review the proposal.\n- Apply the approved change.'}],
               'figures': [{'after_role': role, 'heading': role, 'caption': 'Conceptual ' + role}
                           for role in ('overview', 'workflow')]}
    asset = write_google_doc_source(content, tmp_path)
    document = Document(tmp_path / asset['filename'])
    assert len(document.inline_shapes) == 2
    assert any('image' in rel.reltype for rel in document.sections[0].header.part.rels.values())
    assert 'Overview' not in [p.text for p in document.paragraphs]
    assert [p.text for p in document.paragraphs if p.style.name == 'List Bullet'] == ['Review the proposal.', 'Apply the approved change.']
    assert [p.text for p in document.paragraphs if p.style.name == 'List Number'] == ['Open a run.', 'Review findings.']
    for name in ('Normal', 'Title', 'Heading 1', 'List Number', 'Caption'):
        fonts = document.styles[name].element.rPr.rFonts
        assert fonts.get(qn('w:ascii')) == fonts.get(qn('w:hAnsi')) == 'Nunito'
        assert not any('theme' in attribute.lower() for attribute in fonts.attrib)
    assert document.styles['Title'].element.pPr.find(qn('w:pBdr')) is None


def test_brief_keeps_one_or_two_honest_visuals_and_three_distinct_reasons():
    from app.services.marketing_content import validate_format
    diagram = visual_brief()["visuals"][0]
    with pytest.raises(ValueError, match="at least 1|too_short"):
        validate_format("feature_brief", visual_brief(visuals=[]))
    with pytest.raises(ValueError, match="at most one diagram"):
        validate_format("feature_brief", visual_brief(visuals=[diagram, {**diagram, "heading": "A second diagram"}]))
    with pytest.raises(ValueError, match="verified screenshot id"):
        validate_format("feature_brief", visual_brief(visuals=[{"kind": "screenshot", "heading": "Every channel at once",
                                                                "caption": "Channel Breakdown on a campaign."}]))
    with pytest.raises(ValueError, match="at least three"):
        validate_format("feature_brief", visual_brief(visuals=[{**diagram, "diagram": {**diagram["diagram"], "nodes": diagram["diagram"]["nodes"][:2]}}]))
    reasons = visual_brief()["value"]
    with pytest.raises(ValueError, match="different point"):
        validate_format("feature_brief", visual_brief(value=[*reasons[:2], {**reasons[2], "title": reasons[0]["title"]}]))


def test_carousel_slides_carry_a_visual_not_a_paragraph():
    from app.services.marketing_content import validate_format
    value = validate_format("linkedin_carousel", visual_carousel())
    assert [slide["visual"] for slide in value["slides"]][:3] == ["statement", "contrast", "steps"]
    with pytest.raises(ValueError, match="supporting line|at most"):
        validate_format("linkedin_carousel", visual_carousel(body="A campaign total tells you how much happened. " * 3))
    with pytest.raises(ValueError, match="short labels"):
        validate_format("linkedin_carousel", visual_carousel(points=["A blended total tells you how much happened overall", "Channels"]))
    with pytest.raises(ValueError, match="supported figure"):
        validate_format("linkedin_carousel", visual_carousel(visual="stat", points=["One blended total", "Channel results"]))


def test_carousel_rejects_one_treatment_repeated():
    from app.services.marketing_content import validate_format
    content = visual_carousel()
    content["slides"] = [{**slide, "visual": "statement", "points": []} if slide["layout"] not in {"steps", "problem"}
                         else {**slide, "visual": "steps"} for slide in content["slides"]]
    with pytest.raises(ValueError, match="three distinct visual treatments"):
        validate_format("linkedin_carousel", content)


QUALITY_POST = """A campaign total tells you how much happened. It does not tell you which channel deserves a closer look.

Voice lets you view a breakdown of campaign results by channel. For a sales operations manager preparing a campaign review, that gives the discussion a specific starting point: compare the channel results before deciding which questions to investigate.

For example, imagine reviewing a campaign that uses more than one channel. Start with the breakdown, note where the results differ, and bring those differences to your team. Treat that as a question to explore, not proof that one channel caused a better outcome. Audience, timing and campaign setup may also matter.

The useful shift is from discussing one overall result to asking a more focused question about each channel. What would you investigate first in your next campaign review?

#SalesOperations #CampaignReporting"""


def quality_review(value):
    return {"passed": True, "issues": [], "scores": dict.fromkeys(
        ("factuality", "specificity", "channel_fit", "narrative", "readability"), 5),
        "feature_specificity": "The copy explains a channel breakdown and how a sales operations manager would use that breakdown to structure a review.",
        "audience_value": "Sales operations can ask focused questions without claiming causal attribution from descriptive reporting.",
        "distinct_from_other_formats": "This post opens a focused conversation; there are no other supplied formats with which it overlaps.",
        "claims": [{"location": "body", "content_quote": "Voice lets you view a breakdown of campaign results by channel.",
                    "evidence_quote": QUOTE, "explanation": "The supplied description explicitly supports viewing campaign results broken down by channel."}]}


def test_discovery_repairs_bad_evidence_in_same_window(db, monkeypatch):
    seed(db)
    good = discovered_item()
    bad = {**good, "evidence_quote": "This is a fabricated quote that must never pass."}
    mock_release(monkeypatch, {})
    calls = []
    def answer(surface, prompt, payload, **kwargs):
        calls.append(payload)
        return {"features": [bad if len(calls) == 1 else good]}
    monkeypatch.setattr(manager, "chat_json", answer)
    assert manager.discover_released_features(db)["ok"]
    assert len(calls) == 2 and "verifiable" in calls[1]["repair"]
    assert db.query(MarketingFeature).count() == 2


@pytest.mark.parametrize("verdict", ["skip", "defer"])
def test_rejected_discoveries_are_audited_not_added(db, monkeypatch, verdict):
    seed(db)
    mock_release(monkeypatch, {"features": [discovered_item(verdict)]})
    result = manager.discover_released_features(db)
    assert result["ok"] and result["count"] == 0 and result["rejected"] == 1
    assert db.query(MarketingFeature).count() == 1
    assert list((marketing.CAMPAIGN_DIR / "discovery").glob("*.json"))


def test_bad_window_does_not_block_later_windows(db, monkeypatch):
    seed(db)
    mock_release(monkeypatch, {})
    monkeypatch.setattr(manager, "release_batches", lambda release: [[{"bad": True}], [{"bad": False}]])
    def assess(release, batch, existing):
        if batch[0]["bad"]: raise ValueError("Bad evidence")
        return []
    monkeypatch.setattr(manager, "assess_release_batch", assess)
    result = manager.discover_released_features(db)
    assert not result["ok"] and result["scanned"] == 1 and result["failed_windows"] == 1
    assert db.query(MarketingScan).count() == 1


@pytest.mark.parametrize("path", ["billing-service/docs/plan.md", "billing-service/internal/service/x_test.go",
    "billing-service/static/src/a.test.tsx", "billing-service/secrets.go", "billing-service/scripts/a.py", "another/product.go"])
def test_non_product_source_is_excluded(path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "repo_path_scopes", "billing-service")
    assert not manager.marketing_source(path)
    assert manager.marketing_source("billing-service/internal/campaigns/rotation.go")


def test_legacy_rejected_rows_remain_available_but_not_new(db):
    row = seed(db, decision=decision("skip"), source="codebase", tag="New")
    assert marketing.feature_out(row)["tag"] == "Not selected"
    assert row.tag == "New"  # Historical data remains unchanged.


def test_editorial_quotes_match_across_whitespace_and_glyphs():
    from app.services.marketing_content import validate_review
    from app.services.prompts import parse_json_object
    value = {"title": "Campaign results", "body": QUALITY_POST}
    good = quality_review(value)
    good["claims"][0]["content_quote"] = "Voice lets you view a breakdown of campaign results by channel."
    good["claims"][0]["content_quote"] = good["claims"][0]["content_quote"].replace(" ", "  ")
    assert validate_review(good, value, {"description": QUOTE})["passed"]
    fenced = parse_json_object('```json\n{"passed": true,}\n```')
    assert fenced == {"passed": True}


def test_editorial_scores_do_not_replace_evidence_or_full_coverage():
    from app.services.marketing_content import validate_review
    value = {"title": "Campaign results", "body": QUALITY_POST}
    good = quality_review(value)
    assert validate_review(good, value, {"description": QUOTE})["passed"]
    bad = {**good, "claims": [{**good["claims"][0], "evidence_quote": "Invented supporting evidence that is not in the source."}]}
    with pytest.raises(ValueError, match="supporting source"): validate_review(bad, value, {"description": QUOTE})
    bad = {**good, "issues": ["An unsupported promise remains."]}
    with pytest.raises(ValueError, match="unsupported promise"): validate_review(bad, value, {"description": QUOTE})
    with pytest.raises(ValueError, match="omitted content units"):
        validate_review(good, {**value, "slides": [{"headline": "A new claim", "body": "This also needs an explicit review."}]}, {"description": QUOTE})


def test_editor_cannot_use_generated_positioning_as_proof():
    from app.services.marketing_content import validate_review
    value = {"title": "Campaign results", "body": QUALITY_POST}
    with pytest.raises(ValueError, match="supporting source"):
        validate_review(quality_review(value), value, {"source": "codebase", "description": QUOTE, "benefit": QUOTE})


def test_refresh_reports_marketing_failure(db, monkeypatch, tmp_path):
    from app.services import refresh
    monkeypatch.setattr(refresh, "run_pipeline", lambda *a, **k: SimpleNamespace(status="completed", error="", jira_count=1, cliq_count=1))
    monkeypatch.setattr(refresh, "sync_roadmap_statuses", lambda *a, **k: None)
    monkeypatch.setattr(refresh, "codebase_root", lambda: tmp_path)
    monkeypatch.setattr(refresh, "latest_snapshot", lambda db: None)
    monkeypatch.setattr(refresh, "update_index", lambda *a, **k: {})
    monkeypatch.setattr(refresh, "list_recent_branches", lambda **k: {"branches": []})
    monkeypatch.setattr("app.services.release_worker.detect_and_enqueue", lambda *a, **k: [])
    monkeypatch.setattr("app.services.competitive.refresh_all", lambda *a, **k: {"ok": True, "created": 0})
    monkeypatch.setattr(manager, "refresh_marketing", lambda *a, **k: {"ok": False, "error": "Evidence needs repair"})
    result = refresh.refresh_platform(db)
    assert not result["ok"] and "marketing: Evidence needs repair" in result["error"]
    assert result["steps"]["jira"]["ok"]


def test_workspace_lists_campaigns_without_full_copy(client, db):
    feature = seed(db)
    db.add(MarketingCampaign(
        feature_id=feature.id,
        feature_revision=1,
        feature_snapshot={"name": feature.name, "id": feature.id},
        status="completed",
        content={"version": 2, "formats": {"article": {"status": "completed", "content": {"title": "Long", "body": "x" * 5000}}}},
        assets=[{"filename": "article.md", "channel": "article"}],
        error="",
    ))
    db.commit()
    payload = client.get("/marketing").json()
    campaign = payload["campaigns"][0]
    assert campaign["feature_snapshot"]["name"] == "AI campaigns"
    assert campaign["content"] == {"version": 2}
    assert "formats" not in campaign["content"]
    detail = client.get(f"/marketing/campaigns/{campaign['id']}").json()
    assert detail["content"]["formats"]["article"]["content"]["body"].startswith("x")
